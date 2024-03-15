import csv
import json
import pandas as pd
import psycopg2
import glob
import os
import openpyxl


def transform_to_csv(input_file, output_file):
    # Función para transformar el archivo de entrada a CSV
    if input_file.endswith('.json'):
        with open(input_file, 'r') as file_in, open(output_file, 'w', newline='') as file_out:
            data = json.load(file_in)
            writer = csv.writer(file_out)
            writer.writerow(data[0].keys())  # Escribir encabezados
            for row in data:
                writer.writerow(row.values())
    elif input_file.endswith('.xlsx'):
        data = openpyxl.load_workbook(input_file)
        sheet = data.active
        col = csv.writer(open(input_file[:-5] + "_output.csv", 'w', newline=""))

        for r in sheet.rows:
            if r[0].value is not None:
                col.writerow([cell.value for cell in r])
            else:
                continue

    elif input_file.endswith('.csv'):
        # Si ya es un archivo CSV, simplemente lo copiamos
        #    file_out.write(file_in.read())

        with open(input_file, 'r') as file_in, open(output_file, 'w', newline='') as file_out:
            lines = file_in.readlines()
            for line in lines:
                line = line.strip()
                if line:
                    file_out.write(line + '\n')

    elif input_file.endswith('.db'):
        # Si es una base de datos, no hacemos nada aquí
        pass
    else:
        print("Formato de archivo no compatible:", input_file)


def insert_into_database(csv_file, db_config):
    conn = psycopg2.connect(**db_config)
    cursor = conn.cursor()

    with open(csv_file, 'r', encoding='utf-8') as file:
        reader = csv.reader(file)
        header_row = next(reader)  # Obtener la fila de encabezado
        n_columns = len(header_row)
        column_names = []
        for atr in header_row:
            column_names.append(atr)

        table_name = os.path.splitext(os.path.basename(csv_file))[0]

        # Cambiar el tipo de dato de la primera columna a SERIAL
        columns_def = ", ".join(

            [

                f"{column_names[0]} SERIAL PRIMARY KEY"
                if i == 0
                else f"{column_names[i]} VARCHAR(50)"
                for i in range(n_columns)

            ]
        )




        create_table_query = """
                CREATE TABLE IF NOT EXISTS "{}" ({})
          """.format(table_name, columns_def)
        cursor.execute(create_table_query)

        placeholders = ", ".join(["%s"] * n_columns)
        insert_query = """
          INSERT INTO "{}" ({}) VALUES ({})
          ON CONFLICT DO NOTHING
        """.format(table_name, ", ".join(column_names), placeholders)

        for row in reader:
            cursor.execute(insert_query, row)

    conn.commit()
    conn.close()


def main():
    # input_files = glob.glob('*.csv')  # Obtener todos los archivos CSV en el directorio actual
    current_dir = os.getcwd()

    file_patterns = ["*.csv", "*.json", "*.xlsx", "*.db"]

    input_files = []
    for pattern in file_patterns:
        input_files.extend(glob.glob(os.path.join(current_dir, pattern)))

    csv_output_files = [os.path.splitext(f)[0].replace('\n', '') + '_output.csv' for f in
                        input_files]  # Generar nombres de archivo CSV de salida

    db_config = {
        'host': 'localhost',
        'database': 'test3',
        'user': 'user1',
        'password': 'user1'
    }  # Configuración de la conexión a la base de datos PostgreSQL

    for input_file, csv_output_file in zip(input_files, csv_output_files):
        # Transformar el archivo de entrada a CSV
        transform_to_csv(input_file, csv_output_file)

        # Insertar datos en la base de datos desde el archivo CSV
        insert_into_database(csv_output_file, db_config)


if __name__ == "__main__":
    main()
